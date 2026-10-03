"""Document extraction and SQL-authoritative, recoverable index publication."""
import hashlib
import io
import logging
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from contextlib import contextmanager
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from PIL import Image
from sqlalchemy import select, text

from .config import get_settings
from .db import db_session, DocumentRegistry
from .rag import RagService
from .vision import VisionAnalyzer
from .versioning import VersionIdentity, identify_version

log = logging.getLogger(__name__)
IMAGE_TYPES = {'.png', '.jpg', '.jpeg', '.webp'}
SUPPORTED = {'.pdf', '.doc', '.docx', '.txt', '.md', *IMAGE_TYPES}


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as file:
        for block in iter(lambda: file.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def validate_docx(path, max_uncompressed):
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if 'word/document.xml' not in archive.namelist() or '[Content_Types].xml' not in archive.namelist():
                raise ValueError('Invalid DOCX: the Word document is missing')
            if any(member.flag_bits & 1 for member in members):
                raise ValueError('Encrypted DOCX files are not supported')
            if sum(member.file_size for member in members) > max_uncompressed:
                raise ValueError('DOCX too large after decompression')
        DocxDocument(path)  # Reject malformed XML and broken package relationships.
    except (zipfile.BadZipFile, KeyError, OSError) as error:
        raise ValueError('Invalid DOCX document') from error
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('Invalid DOCX document') from error


def validate_pdf(path):
    try:
        with fitz.open(path) as pdf:
            if not pdf.is_pdf:
                raise ValueError('Invalid PDF document')
            if pdf.needs_pass:
                raise ValueError('Password-protected PDFs are not supported')
            if not pdf.page_count:
                raise ValueError('PDF contains no pages')
            if pdf.page_count > getattr(get_settings(), 'admin_pdf_max_pages', 1000):
                raise ValueError('PDF exceeds the configured page limit')
    except (fitz.FileDataError, fitz.EmptyFileError, RuntimeError) as error:
        raise ValueError('Invalid PDF document') from error


def _libreoffice_executable():
    for name in ('libreoffice', 'soffice'):
        executable = shutil.which(name)
        if executable:
            return executable
    for candidate in (
        Path('C:/Program Files/LibreOffice/program/soffice.exe'),
        Path('C:/Program Files (x86)/LibreOffice/program/soffice.exe'),
    ):
        if candidate.is_file():
            return str(candidate)
    return None


def validate_document(path, max_uncompressed=None):
    path = Path(path)
    if not path.is_file() or not path.stat().st_size:
        raise ValueError('Document is empty or missing')
    extension = path.suffix.lower()
    if extension not in SUPPORTED:
        raise ValueError('Unsupported document type')
    if extension == '.docx':
        validate_docx(
            path, max_uncompressed or get_settings().admin_docx_max_uncompressed_bytes,
        )
    elif extension == '.pdf':
        validate_pdf(path)
    elif extension == '.doc':
        with path.open('rb') as file:
            if file.read(8) != b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                raise ValueError('Invalid legacy Word document; save it as DOCX')
        if not _libreoffice_executable():
            raise ValueError('Legacy DOC requires local LibreOffice; save it as DOCX instead')
    elif extension in IMAGE_TYPES:
        try:
            with Image.open(path) as image:
                image.verify()
        except (OSError, ValueError, Image.DecompressionBombError) as error:
            raise ValueError('Invalid image document') from error


@contextmanager
def _readable_path(path):
    if path.suffix.lower() != '.doc':
        yield path
        return
    executable = _libreoffice_executable()
    if not executable:
        raise ValueError('Legacy DOC requires local LibreOffice; save it as DOCX instead')
    # A separate user profile prevents conversions from joining another LO process.
    with tempfile.TemporaryDirectory(prefix='rag-doc-') as folder:
        temporary = Path(folder)
        profile = (temporary / 'profile').as_uri()
        command = [
            executable, f'-env:UserInstallation={profile}', '--headless',
            '--convert-to', 'docx', '--outdir', str(temporary), str(path.resolve()),
        ]
        try:
            result = subprocess.run(
                command, capture_output=True, timeout=getattr(
                    get_settings(), 'legacy_doc_conversion_timeout_seconds', 120,
                ), check=False,
            )
        except subprocess.TimeoutExpired as error:
            raise ValueError('Local Word conversion timed out') from error
        converted = temporary / f'{path.stem}.docx'
        if result.returncode or not converted.is_file():
            raise ValueError('Local LibreOffice could not convert this DOC; save it as DOCX')
        validate_docx(converted, get_settings().admin_docx_max_uncompressed_bytes)
        yield converted


def ocr_image(data):
    settings = get_settings()
    executable = shutil.which('tesseract')
    if not executable:
        raise ValueError('Scanned documents require local Tesseract with eng and heb language data')
    try:
        result = subprocess.run(
            [executable, 'stdin', 'stdout', '-l', getattr(settings, 'ocr_languages', 'eng+heb'), '--psm', '3'],
            input=data, capture_output=True,
            timeout=getattr(settings, 'ocr_timeout_seconds', 120), check=False,
        )
    except subprocess.TimeoutExpired as error:
        raise ValueError('Local OCR timed out; try a smaller document or clearer scan') from error
    if result.returncode:
        raise ValueError('Local OCR failed; verify Tesseract and the configured language data')
    return result.stdout.decode('utf-8', errors='replace').strip()


def _block_text(parent):
    element = parent.element.body if hasattr(parent, 'element') and hasattr(parent.element, 'body') else parent._element
    for child in element.iterchildren():
        kind = child.tag.rsplit('}', 1)[-1]
        if kind == 'p':
            value = Paragraph(child, parent).text.strip()
            if value:
                yield value
        elif kind == 'tbl':
            value = _table_text(Table(child, parent))
            if value:
                yield value


def _table_text(table):
    rows = []
    for row in table.rows:
        seen = set()
        values = []
        for cell in row.cells:
            # Merged cells can appear several times in python-docx's grid.
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            values.append('\n'.join(_block_text(cell)))
        if any(values):
            rows.append(' | '.join(values))
    return '\n'.join(rows)


def _word_text(path):
    document = DocxDocument(path)
    headers, footers, seen = [], [], set()
    for section in document.sections:
        for label, destination, areas in (
            ('Header', headers, (section.header, section.first_page_header, section.even_page_header)),
            ('Footer', footers, (section.footer, section.first_page_footer, section.even_page_footer)),
        ):
            for area in areas:
                part = str(area.part.partname)
                if part in seen:
                    continue
                seen.add(part)
                value = '\n'.join(_block_text(area))
                if value:
                    destination.append(f'{label}:\n{value}')
    return '\n\n'.join([*headers, *_block_text(document), *footers])


def text_docs(path):
    path = Path(path)
    settings = get_settings()
    if path.suffix.lower() == '.doc':
        with _readable_path(path) as converted:
            yield from text_docs(converted)
    elif path.suffix.lower() == '.pdf':
        with fitz.open(path) as pdf:
            for index, page in enumerate(pdf):
                value = page.get_text('text', sort=True).strip()
                extracted = 'text'
                # Keep a usable text layer. Render scans as a whole page so OCR
                # preserves layout instead of indexing detached image fragments.
                needs_ocr = len(value) < getattr(settings, 'ocr_min_text_chars', 24) and bool(page.get_images(full=True))
                if needs_ocr and getattr(settings, 'ocr_enabled', True):
                    scale = getattr(settings, 'ocr_pdf_dpi', 200) / 72
                    largest = max(page.rect.width, page.rect.height)
                    if largest * scale > 4000:
                        scale = 4000 / largest
                    pixels = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                    recognized = ocr_image(pixels.tobytes('png'))
                    if recognized:
                        value = recognized
                        extracted = 'ocr'
                if value:
                    yield Document(page_content=value, metadata={
                        'page': index + 1, 'content_type': extracted,
                    })
    elif path.suffix.lower() == '.docx':
        value = _word_text(path)
        if value:
            yield Document(page_content=value, metadata={'content_type': 'text'})
    elif path.suffix.lower() not in IMAGE_TYPES:
        value = path.read_text(encoding='utf-8', errors='replace')
        if value.strip():
            yield Document(page_content=value, metadata={'content_type': 'text'})


def images(path):
    path = Path(path)
    if path.suffix.lower() in IMAGE_TYPES:
        yield path.read_bytes(), None, path.name
    elif path.suffix.lower() == '.docx':
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith('word/media/') and not name.endswith('/'):
                    yield archive.read(name), None, name
    elif path.suffix.lower() == '.pdf':
        with fitz.open(path) as pdf:
            for page_number, page in enumerate(pdf, 1):
                for image_number, image in enumerate(page.get_images(full=True), 1):
                    try:
                        data = pdf.extract_image(image[0]).get('image', b'')
                    except (ValueError, RuntimeError):
                        continue
                    if data:
                        yield data, page_number, f'page {page_number} image {image_number}'


def topic_lock(db, topic):
    if db.get_bind().dialect.name == 'postgresql':
        lock_id = int.from_bytes(hashlib.sha256(topic.encode()).digest()[:8], 'big', signed=True)
        db.execute(text('SELECT pg_advisory_xact_lock(:lock_id)'), {'lock_id': lock_id})


def _active_rows(db, topic):
    return db.scalars(select(DocumentRegistry).where(
        DocumentRegistry.topic_key == topic, DocumentRegistry.status == 'active',
    ).order_by(DocumentRegistry.version_rank.desc(), DocumentRegistry.id.desc())).all()


def _skip_result(active, identity, digest):
    if active and identity.version_rank < active.version_rank:
        return {
            'status': 'skipped_older', 'active_version': active.version_label,
            'topic': identity.topic_key, 'generation_id': active.generation_id,
        }
    if active and active.sha256 == digest and identity.version_rank == active.version_rank:
        return {
            'status': 'unchanged', 'active_version': active.version_label,
            'topic': identity.topic_key, 'generation_id': active.generation_id,
        }
    return None


def ingest_path(path_str, origin='admin', *, source_name=None, topic_key=None, version_rank=None, version_label=None, publication_guard=None, persist_result=None):
    settings = get_settings()
    path = Path(path_str)
    validate_document(path, settings.admin_docx_max_uncompressed_bytes)
    source = source_name or path.name
    derived = identify_version(source, origin)
    identity = VersionIdentity(
        topic_key or derived.topic_key, derived.topic_name,
        derived.version_rank if version_rank is None else int(version_rank),
        version_label or (f'v{version_rank}' if version_rank else derived.version_label),
    )
    if not identity.topic_key or len(identity.topic_key) > 255:
        raise ValueError('Document topic must contain between 1 and 255 characters')
    if identity.version_rank < 0 or identity.version_rank > 2147483647 or len(identity.version_label) > 80:
        raise ValueError('Document version is outside the supported range')
    digest = sha256_file(path)
    with db_session() as db:
        if publication_guard:
            publication_guard(db)
        existing = _active_rows(db, identity.topic_key)
        skipped = _skip_result(existing[0] if existing else None, identity, digest)
        if skipped:
            if persist_result:
                persist_result(db, skipped)
            return skipped
    generation = uuid.uuid4().hex
    rag = RagService()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.ingest_chunk_size, chunk_overlap=settings.ingest_chunk_overlap,
    )
    documents = []
    with _readable_path(path) as readable:
        bases = list(text_docs(readable))
        analyzer = VisionAnalyzer(rag.client) if settings.vision_enabled else None
        if analyzer or (getattr(settings, 'ocr_enabled', True) and readable.suffix.lower() != '.pdf'):
            for index, (data, page, label) in enumerate(images(readable)):
                if index >= settings.vision_max_images_per_document:
                    break
                # OCR also works without an optional vision model. Ignore vector
                # Word art, but do not silently ignore a missing OCR installation.
                if getattr(settings, 'ocr_enabled', True) and readable.suffix.lower() != '.pdf':
                    try:
                        with Image.open(io.BytesIO(data)) as image:
                            image.load()
                            is_raster = True
                    except (OSError, ValueError, Image.DecompressionBombError):
                        is_raster = False
                    if is_raster:
                        recognized = ocr_image(data)
                        if recognized:
                            metadata = {'content_type': 'ocr'}
                            if page is not None:
                                metadata['page'] = page
                            bases.append(Document(page_content=f'Image text ({label}):\n{recognized}', metadata=metadata))
                if analyzer:
                    value = analyzer.analyze(data)
                    if value:
                        metadata = {'content_type': 'image'}
                        if page is not None:
                            metadata['page'] = page
                        bases.append(Document(page_content=f'Visual evidence ({label}):\n{value}', metadata=metadata))
        for base in bases:
            for chunk in splitter.split_documents([base]):
                chunk.metadata.update({
                    'source': source, 'origin': origin, 'topic_key': identity.topic_key,
                    'version': identity.version_label, 'generation_id': generation,
                    'active': False, 'chunk_index': len(documents),
                })
                documents.append(chunk)
    if not documents:
        raise ValueError('No readable text found; use a clear scan, a PDF text layer, or a Word document containing text')
    ids = [hashlib.sha256(f'{generation}|{index}|{chunk.page_content}'.encode()).hexdigest()
           for index, chunk in enumerate(documents)]
    with db_session() as db:
        if publication_guard:
            publication_guard(db)
        db.add(DocumentRegistry(
            topic_key=identity.topic_key, source=source, storage_path=str(path),
            origin=origin, version_label=identity.version_label,
            version_rank=identity.version_rank, sha256=digest, generation_id=generation,
            status='staging', chunk_count=len(documents),
        ))
    old_generations = []
    try:
        rag.index_documents(documents, ids)
        # Indexing completes before a short, serialized publication transaction.
        # Retrieval queries the SQL registry, so no Chroma flag update is required.
        with db_session() as db:
            topic_lock(db, identity.topic_key)
            if publication_guard:
                publication_guard(db)
            previous = _active_rows(db, identity.topic_key)
            current = db.scalar(select(DocumentRegistry).where(DocumentRegistry.generation_id == generation))
            if current is None or current.status != 'staging':
                raise ValueError('Document indexing was cancelled before publication')
            # A reclaimed job may have left an incomplete generation behind.
            # Immutable upload paths identify retries without cancelling another
            # upload of a different version or different document content.
            abandoned = db.scalars(select(DocumentRegistry).where(
                DocumentRegistry.topic_key == identity.topic_key,
                DocumentRegistry.storage_path == str(path),
                DocumentRegistry.sha256 == digest,
                DocumentRegistry.version_rank == identity.version_rank,
                DocumentRegistry.status == 'staging',
                DocumentRegistry.generation_id != generation,
            )).all()
            for row in abandoned:
                row.status = 'superseded'
                old_generations.append(row.generation_id)
            skipped = _skip_result(previous[0] if previous else None, identity, digest)
            if skipped:
                current.status = 'superseded'
            else:
                for row in previous:
                    row.status = 'superseded'
                    old_generations.append(row.generation_id)
                current.status = 'active'
            published_result = ({**skipped, 'cleanup_pending': old_generations + [generation]} if skipped else {
                'status': 'indexed', 'topic': identity.topic_key, 'version': identity.version_label,
                'chunks': len(documents), 'generation_id': generation,
                'cleanup_pending': list(old_generations),
            })
            if persist_result:
                persist_result(db, published_result)
    except Exception:
        # A connection loss at COMMIT can leave its outcome uncertain. Never
        # delete vectors until SQL confirms they were not published.
        can_delete = False
        try:
            with db_session() as db:
                failed = db.scalar(select(DocumentRegistry).where(DocumentRegistry.generation_id == generation))
                if failed and failed.status != 'active':
                    if failed.status == 'staging':
                        failed.status = 'failed'
                    can_delete = True
        except Exception:
            log.exception('Could not confirm publication status for generation %s; preserving vectors', generation)
        if can_delete:
            try:
                rag.delete_generation(generation)
            except Exception:
                log.exception('Failed to clean staged generation %s; it remains excluded from retrieval', generation)
        raise
    if skipped:
        old_generations.append(generation)
    cleanup_pending = []
    for old_generation in old_generations:
        try:
            rag.delete_generation(old_generation)
        except Exception:
            cleanup_pending.append(old_generation)
            log.exception('Generation %s is retired in SQL but awaits vector cleanup', old_generation)
    if skipped:
        return {**skipped, 'cleanup_pending': cleanup_pending}
    return {
        'status': 'indexed', 'topic': identity.topic_key, 'version': identity.version_label,
        'chunks': len(documents), 'generation_id': generation,
        'cleanup_pending': cleanup_pending,
    }


def repair_legacy_documents():
    """Reindex timestamp-staged uploads; retire each old topic only after success.

    This is an explicit maintenance operation. It preserves the old active
    generation whenever a source is missing or indexing cannot finish.
    """
    with db_session() as db:
        records = list(db.scalars(select(DocumentRegistry).where(DocumentRegistry.status == 'active')).all())
    results = []
    for record in records:
        storage = Path(record.storage_path or record.source)
        match = re.match(r'^\d{10,}_(.+)$', storage.name)
        if not match or record.origin != 'admin':
            continue
        original_name = match.group(1)
        stable_topic = identify_version(original_name, 'admin').topic_key
        if record.topic_key == stable_topic:
            continue
        try:
            result = ingest_path(
                str(storage), 'admin', source_name=original_name, topic_key=stable_topic,
                version_rank=record.version_rank, version_label=record.version_label,
            )
            retired = False
            with db_session() as db:
                topic_lock(db, stable_topic)
                replacement = _active_rows(db, stable_topic)
                old = db.get(DocumentRegistry, record.id)
                if replacement and old and old.status == 'active' and old.generation_id != replacement[0].generation_id:
                    old.status = 'superseded'
                    retired = True
            if retired:
                try:
                    RagService().delete_generation(record.generation_id)
                except Exception:
                    log.exception('Legacy generation %s awaits vector cleanup', record.generation_id)
            results.append({'source': original_name, 'status': result['status'], 'retired_legacy': retired})
        except Exception as error:
            results.append({'source': original_name, 'status': 'failed', 'error': str(error)})
    return {'documents': results, 'repaired': sum(row.get('retired_legacy', False) for row in results)}
