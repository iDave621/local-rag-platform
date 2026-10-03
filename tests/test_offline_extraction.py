"""Actual local OCR/LibreOffice acceptance, intended for the release API image."""
import io
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))

import fitz
from docx import Document as WordDocument
from PIL import Image, ImageDraw, ImageFont, features

from app import ingestion
from app.config import Settings


class OfflineExtractionAcceptance(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='rag-offline-')
        self.root = Path(self.temporary.name)
        self.settings = Settings(_env_file=None, database_url='sqlite://', ocr_enabled=True)
        self.config = patch.object(ingestion, 'get_settings', return_value=self.settings)
        self.config.start()

    def tearDown(self):
        self.config.stop()
        self.temporary.cleanup()

    def raster(self, text, hebrew=False):
        font_path = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
        if not Path(font_path).exists() or not shutil.which('tesseract'):
            self.skipTest('Release image with Tesseract and DejaVu fonts required')
        image = Image.new('RGB', (1200, 240), 'white')
        draw = ImageDraw.Draw(image)
        if hebrew and not features.check('raqm'):
            font = ImageFont.truetype(font_path, 55, layout_engine=ImageFont.Layout.BASIC)
            draw.text((60, 75), text[::-1], fill='black', font=font)
        else:
            font = ImageFont.truetype(font_path, 55)
            draw.text((60, 75), text, fill='black', font=font, **({'direction': 'rtl'} if hebrew else {}))
        data = io.BytesIO()
        image.save(data, 'PNG')
        return data.getvalue()

    def test_actual_english_scanned_pdf_with_network_disabled(self):
        pixels = self.raster('Approved VLAN 120 port 7')
        path = self.root / 'scan.pdf'
        with fitz.open() as pdf:
            page = pdf.new_page(width=600, height=120)
            page.insert_image(page.rect, stream=pixels)
            pdf.save(path)
        ingestion.validate_document(path)
        documents = list(ingestion.text_docs(path))
        self.assertEqual(len(documents), 1)
        self.assertIn('120', documents[0].page_content)
        self.assertIn('VLAN', documents[0].page_content)
        self.assertEqual(documents[0].metadata['content_type'], 'ocr')
        print('Actual PDF OCR:', documents[0].page_content)

    def test_actual_hebrew_raster_with_network_disabled(self):
        value = ingestion.ocr_image(self.raster('הגדרות רשת מקומית', hebrew=True))
        self.assertIn('רשת', value)
        self.assertIn('מקומית', value)
        print('Actual Hebrew OCR:', value)

    def test_actual_legacy_word_conversion_with_network_disabled(self):
        executable = ingestion._libreoffice_executable()
        if not executable:
            self.skipTest('Release image with LibreOffice required')
        source = self.root / 'legacy.docx'
        document = WordDocument()
        document.add_paragraph('Local legacy Word evidence VLAN 120')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Port'
        table.cell(0, 1).text = '7'
        document.add_paragraph('End of local procedure')
        document.save(source)
        output = self.root / 'old-format'
        output.mkdir()
        converted = subprocess.run([
            executable, f'-env:UserInstallation={(self.root / "creation-profile").as_uri()}',
            '--headless', '--convert-to', 'doc:MS Word 97',
            '--outdir', str(output), str(source),
        ], capture_output=True, timeout=120, check=False)
        self.assertEqual(converted.returncode, 0, converted.stderr.decode(errors='replace'))
        legacy = output / 'legacy.doc'
        self.assertTrue(legacy.is_file())
        ingestion.validate_document(legacy)
        extracted = '\n'.join(doc.page_content for doc in ingestion.text_docs(legacy))
        self.assertIn('VLAN 120', extracted)
        self.assertIn('Port | 7', extracted)
        print('Actual legacy DOC extraction:', extracted)


if __name__ == '__main__':
    unittest.main()
