"""Offline regressions for extraction, publication, grounding and sync retries.

Run with the backend dependencies installed:
  python -m unittest discover -s tests -p test_ingestion_rag.py -v
All databases and documents are temporary; no Ollama/network calls are made.
"""
import io
import json
import sys
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))

import fitz
import httpx
from docx import Document as WordDocument
from langchain_core.documents import Document
from PIL import Image
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app import ingestion, literature, rag, vision
from app.config import Settings
from app.db import Base, DocumentRegistry, JobRecord
from app.queueing import JobQueue
from app.versioning import identify_version


class TemporaryStore(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='rag-tests-')
        self.root = Path(self.temporary.name)
        self.settings = Settings(
            _env_file=None, database_url='sqlite://', data_root=str(self.root),
            ocr_enabled=False, vision_enabled=False,
        )
        self.engine = create_engine('sqlite://')
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(self.engine, expire_on_commit=False)

        @contextmanager
        def database():
            session = self.sessions()
            try:
                yield session
                session.commit()
            except Exception:
                session.rollback()
                raise
            finally:
                session.close()

        self.database = database
        self.patches = [
            patch.object(ingestion, 'get_settings', return_value=self.settings),
            patch.object(ingestion, 'db_session', database),
            patch.object(rag, 'db_session', database),
            patch.object(literature, 'db_session', database),
        ]
        for handle in self.patches:
            handle.start()

    def tearDown(self):
        for handle in reversed(self.patches):
            handle.stop()
        self.engine.dispose()
        self.temporary.cleanup()

    def word(self, name, value):
        path = self.root / name
        document = WordDocument()
        document.add_paragraph(value)
        document.save(path)
        return path

    def active(self):
        with self.database() as db:
            return list(db.scalars(select(DocumentRegistry).where(DocumentRegistry.status == 'active')).all())


class MemoryVectors:
    def __init__(self):
        self.client = MagicMock()
        self.documents = {}
        self.calls = []
        self.fail_index = False
        self.fail_delete = False
        self.after_index = None

    def index_documents(self, documents, ids):
        if len(ids) != len(set(ids)):
            raise AssertionError('Duplicate chunk IDs')
        self.calls.append((list(documents), list(ids)))
        self.documents.update(zip(ids, documents))
        if self.fail_index:
            raise RuntimeError('Injected embedding failure')
        if self.after_index:
            self.after_index(documents[0].metadata['generation_id'])

    def delete_generation(self, generation):
        if self.fail_delete:
            raise RuntimeError('Injected vector deletion failure')
        self.documents = {
            key: value for key, value in self.documents.items()
            if value.metadata['generation_id'] != generation
        }


class ExtractionTests(TemporaryStore):
    def test_word_body_tables_nested_headers_and_footers_in_order(self):
        path = self.root / 'manual.docx'
        document = WordDocument()
        document.sections[0].header.paragraphs[0].text = 'Equipment header'
        document.add_paragraph('Before configuration')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'VLAN'
        table.cell(0, 1).text = '120'
        nested = table.cell(0, 1).add_table(rows=1, cols=2)
        nested.cell(0, 0).text = 'Port'
        nested.cell(0, 1).text = '7'
        document.add_paragraph('After configuration')
        document.sections[0].footer.paragraphs[0].text = 'Revision footer'
        document.save(path)
        extracted = list(ingestion.text_docs(path))[0].page_content
        for value in ('Equipment header', 'VLAN | 120', 'Port | 7', 'Revision footer'):
            self.assertIn(value, extracted)
        self.assertLess(extracted.index('Before configuration'), extracted.index('VLAN'))
        self.assertLess(extracted.index('Port | 7'), extracted.index('After configuration'))

    def test_table_only_word_is_indexable(self):
        path = self.root / 'table.docx'
        document = WordDocument()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Gateway'
        table.cell(0, 1).text = '10.0.0.1'
        document.save(path)
        self.assertIn('Gateway | 10.0.0.1', list(ingestion.text_docs(path))[0].page_content)

    def test_scanned_pdf_ocr_is_page_local(self):
        self.settings.ocr_enabled = True
        path = self.root / 'scanned.pdf'
        image = io.BytesIO()
        Image.new('RGB', (300, 200), 'white').save(image, 'PNG')
        with fitz.open() as pdf:
            pdf.new_page().insert_text((20, 40), 'Digital page contains sufficient usable text')
            page = pdf.new_page()
            page.insert_image(page.rect, stream=image.getvalue())
            pdf.save(path)
        with patch.object(ingestion, 'ocr_image', return_value='Recovered scan: VLAN 120') as ocr:
            documents = list(ingestion.text_docs(path))
        self.assertEqual(ocr.call_count, 1)
        self.assertEqual([doc.metadata['page'] for doc in documents], [1, 2])
        self.assertEqual(documents[1].metadata['content_type'], 'ocr')
        self.assertIn('VLAN 120', documents[1].page_content)

    def test_encrypted_pdf_rejected_before_queueing(self):
        path = self.root / 'locked.pdf'
        with fitz.open() as pdf:
            pdf.new_page()
            pdf.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='secret', owner_pw='owner')
        with self.assertRaisesRegex(ValueError, 'Password'):
            ingestion.validate_document(path)

    def test_invalid_pdf_and_docx_rejected(self):
        for name in ('bad.pdf', 'bad.docx'):
            path = self.root / name
            path.write_bytes(b'not a document')
            with self.assertRaises(ValueError):
                ingestion.validate_document(path)

    def test_legacy_doc_requires_local_converter(self):
        path = self.root / 'old.doc'
        path.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'old Word')
        with patch.object(ingestion, '_libreoffice_executable', return_value=None):
            with self.assertRaisesRegex(ValueError, 'LibreOffice'):
                ingestion.validate_document(path)

    def test_legacy_conversion_extracts_local_word_output(self):
        path = self.root / 'old.doc'
        path.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'old Word')

        def convert(command, **kwargs):
            destination = Path(command[command.index('--outdir') + 1]) / 'old.docx'
            document = WordDocument()
            document.add_paragraph('Converted legacy procedure')
            document.save(destination)
            return SimpleNamespace(returncode=0)

        with patch.object(ingestion, '_libreoffice_executable', return_value='local-soffice'):
            with patch.object(ingestion.subprocess, 'run', side_effect=convert) as run:
                documents = list(ingestion.text_docs(path))
        self.assertIn('Converted legacy procedure', documents[0].page_content)
        self.assertNotIn('shell', run.call_args.kwargs)

    def test_ocr_does_not_download_models(self):
        self.settings.ocr_enabled = True
        with patch.object(ingestion.shutil, 'which', return_value='local-tesseract'):
            with patch.object(ingestion.subprocess, 'run', return_value=SimpleNamespace(
                returncode=0, stdout='תקלה 120'.encode(), stderr=b'',
            )) as run:
                self.assertEqual(ingestion.ocr_image(b'local pixels'), 'תקלה 120')
        command = run.call_args.args[0]
        self.assertEqual(command[command.index('-l') + 1], 'eng+heb')
        self.assertNotIn('shell', run.call_args.kwargs)


class PublicationTests(TemporaryStore):
    def setUp(self):
        super().setUp()
        self.vectors = MemoryVectors()
        self.rag_patch = patch.object(ingestion, 'RagService', return_value=self.vectors)
        self.rag_patch.start()

    def tearDown(self):
        self.rag_patch.stop()
        super().tearDown()

    def test_duplicate_pdf_pages_use_distinct_chunk_ids(self):
        path = self.root / 'same.pdf'
        with fitz.open() as pdf:
            for _ in range(2):
                pdf.new_page().insert_text((30, 40), 'Identical repeated instructions for VLAN 120')
            pdf.save(path)
        result = ingestion.ingest_path(str(path))
        self.assertEqual(result['chunks'], 2)
        ids = self.vectors.calls[0][1]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual([doc.metadata['page'] for doc in self.vectors.calls[0][0]], [1, 2])

    def test_original_identity_survives_unique_storage_paths(self):
        older = self.word('uuid-one.docx', 'Old configuration: VLAN 100')
        newer = self.word('uuid-two.docx', 'New configuration: VLAN 120')
        first = ingestion.ingest_path(str(older), source_name='guide_v1.docx')
        second = ingestion.ingest_path(str(newer), source_name='guide_v2.docx')
        self.assertEqual(first['topic'], second['topic'])
        self.assertEqual(self.active()[0].source, 'guide_v2.docx')
        self.assertEqual(self.active()[0].version_rank, 2)
        self.assertEqual(self.active()[0].storage_path, str(newer))
        self.assertEqual(len(self.active()), 1)

    def durable_job(self, path, jid='publication-job'):
        with self.database() as db:
            db.add(JobRecord(id=jid, queue='medium', kind='ingest',
                             payload={'path': str(path)}, status='processing',
                             claim_token='publication-lease', attempts=1, max_attempts=1))
        queue = JobQueue()
        job = {'id': jid, 'claim_token': 'publication-lease'}
        return lambda db, result: queue.persist_result(db, job, result)

    def test_publication_saves_its_job_result_in_the_registry_transaction(self):
        path = self.word('guide.docx', 'Management VLAN 120')
        result = ingestion.ingest_path(str(path), persist_result=self.durable_job(path))
        with self.database() as db:
            row = db.get(JobRecord, 'publication-job')
            self.assertEqual(row.result['generation_id'], result['generation_id'])
            self.assertEqual(row.result['status'], 'indexed')
            self.assertEqual(row.status, 'processing')
        self.assertEqual(self.active()[0].generation_id, result['generation_id'])

    def test_result_save_failure_rolls_back_publication_and_keeps_old_generation(self):
        older = self.word('guide_v1.docx', 'Management VLAN 100')
        newer = self.word('guide_v2.docx', 'Management VLAN 120')
        first = ingestion.ingest_path(str(older))
        persist = self.durable_job(newer)

        def fail_result(db, result):
            persist(db, result)
            raise RuntimeError('Injected result transaction failure')

        with self.assertRaisesRegex(RuntimeError, 'result transaction'):
            ingestion.ingest_path(str(newer), persist_result=fail_result)
        self.assertEqual(self.active()[0].generation_id, first['generation_id'])
        with self.database() as db:
            self.assertIsNone(db.get(JobRecord, 'publication-job').result)

    def test_unchanged_document_also_saves_a_durable_job_result(self):
        path = self.word('guide.docx', 'Management VLAN 120')
        first = ingestion.ingest_path(str(path))
        result = ingestion.ingest_path(str(path), persist_result=self.durable_job(path))
        self.assertEqual(result['status'], 'unchanged')
        self.assertEqual(result['generation_id'], first['generation_id'])
        with self.database() as db:
            self.assertEqual(db.get(JobRecord, 'publication-job').result, result)

    def test_failed_replacement_keeps_previous_active_generation_and_recovers(self):
        older = self.word('guide_v1.docx', 'Old procedure')
        newer = self.word('guide_v2.docx', 'Updated procedure')
        first = ingestion.ingest_path(str(older))
        self.vectors.fail_index = True
        with self.assertRaisesRegex(RuntimeError, 'embedding'):
            ingestion.ingest_path(str(newer))
        self.assertEqual(self.active()[0].generation_id, first['generation_id'])
        self.vectors.fail_index = False
        result = ingestion.ingest_path(str(newer))
        self.assertEqual(self.active()[0].generation_id, result['generation_id'])
        self.assertEqual(self.active()[0].version_rank, 2)

    def test_deletion_failure_cannot_keep_retired_generation_retrievable(self):
        older = self.word('guide_v1.docx', 'Old procedure')
        newer = self.word('guide_v2.docx', 'Updated procedure')
        first = ingestion.ingest_path(str(older))
        self.vectors.fail_delete = True
        with self.assertLogs(ingestion.log, level='ERROR'):
            second = ingestion.ingest_path(str(newer))
        self.assertIn(first['generation_id'], second['cleanup_pending'])
        service = rag.RagService.__new__(rag.RagService)
        self.assertEqual(service._active_generations(), {second['generation_id']})

    def test_publication_refuses_cancelled_staging_record(self):
        path = self.word('guide.docx', 'A procedure')

        def cancel(generation):
            with self.database() as db:
                row = db.scalar(select(DocumentRegistry).where(DocumentRegistry.generation_id == generation))
                row.status = 'deleted'

        self.vectors.after_index = cancel
        with self.assertRaisesRegex(ValueError, 'cancelled'):
            ingestion.ingest_path(str(path))
        self.assertEqual(self.active(), [])
        self.assertEqual(self.vectors.documents, {})

    def test_lease_loss_before_publication_keeps_previous_generation(self):
        older = self.word('guide_v1.docx', 'Old procedure')
        newer = self.word('guide_v2.docx', 'New procedure')
        initial = ingestion.ingest_path(str(older))
        calls = []

        def guard(db):
            calls.append(db)
            if len(calls) == 3:
                raise RuntimeError('Job lease was reassigned')

        with self.assertRaisesRegex(RuntimeError, 'reassigned'):
            ingestion.ingest_path(str(newer), publication_guard=guard)
        self.assertEqual(self.active()[0].generation_id, initial['generation_id'])
        self.assertEqual({doc.metadata['generation_id'] for doc in self.vectors.documents.values()},
                         {initial['generation_id']})

    def test_lease_loss_before_staging_creates_no_registry_or_vectors(self):
        path = self.word('guide.docx', 'A procedure')
        calls = []

        def guard(db):
            calls.append(db)
            if len(calls) == 2:
                raise RuntimeError('Job was cancelled during extraction')

        with self.assertRaisesRegex(RuntimeError, 'cancelled'):
            ingestion.ingest_path(str(path), publication_guard=guard)
        with self.database() as db:
            self.assertEqual(list(db.scalars(select(DocumentRegistry)).all()), [])
        self.assertEqual(self.vectors.calls, [])

    def test_retry_publication_retires_interrupted_generation_for_same_upload(self):
        path = self.word('guide_v2.docx', 'Recovered upload')
        with self.database() as db:
            db.add(DocumentRegistry(
                topic_key='admin/guide', source=path.name, storage_path=str(path),
                version_label='v2', version_rank=2, sha256=ingestion.sha256_file(path),
                generation_id='interrupted', status='staging',
            ))
        self.vectors.documents['orphan-vector'] = Document(page_content='Incomplete old chunk', metadata={
            'generation_id': 'interrupted',
        })
        result = ingestion.ingest_path(str(path))
        with self.database() as db:
            abandoned = db.scalar(select(DocumentRegistry).where(DocumentRegistry.generation_id == 'interrupted'))
            self.assertEqual(abandoned.status, 'superseded')
        self.assertEqual(self.active()[0].generation_id, result['generation_id'])
        self.assertNotIn('orphan-vector', self.vectors.documents)

    def test_preexisting_duplicate_active_topics_use_only_latest_version(self):
        with self.database() as db:
            for rank in (1, 3, 2):
                db.add(DocumentRegistry(
                    topic_key='admin/guide', source=f'guide_v{rank}.docx',
                    version_label=f'v{rank}', version_rank=rank, sha256=str(rank) * 64,
                    generation_id=f'generation-{rank}', status='active',
                ))
        service = rag.RagService.__new__(rag.RagService)
        self.assertEqual(service._active_generations(), {'generation-3'})

    def test_older_version_never_overwrites_current(self):
        newer = self.word('guide_v3.docx', 'Current procedure')
        older = self.word('guide_v1.docx', 'Old procedure')
        ingestion.ingest_path(str(newer))
        result = ingestion.ingest_path(str(older))
        self.assertEqual(result['status'], 'skipped_older')
        self.assertEqual(len(self.vectors.calls), 1)
        self.assertEqual(self.active()[0].version_rank, 3)

    def test_supplied_remote_identity_is_preserved(self):
        path = self.word('hashed_manual.docx', 'Manual evidence')
        result = ingestion.ingest_path(
            str(path), 'internal_site', source_name='manual.docx',
            topic_key='internal_site/intranet/routers/manual', version_rank=7, version_label='v7',
        )
        self.assertEqual(result['topic'], 'internal_site/intranet/routers/manual')
        self.assertEqual(self.active()[0].version_rank, 7)
        self.assertEqual(self.vectors.calls[0][0][0].metadata['source'], 'manual.docx')

    def test_legacy_repair_preserves_source_and_is_idempotent(self):
        path = self.word('1770000000_guide_v1.docx', 'Legacy staged configuration')
        with self.database() as db:
            db.add(DocumentRegistry(
                topic_key='admin/1770000000 guide', source=str(path),
                version_label='v1', version_rank=1, sha256=ingestion.sha256_file(path),
                generation_id='legacy', status='active',
            ))
        result = ingestion.repair_legacy_documents()
        self.assertEqual(result['repaired'], 1)
        self.assertEqual(self.active()[0].topic_key, 'admin/guide')
        self.assertEqual(ingestion.repair_legacy_documents()['repaired'], 0)

    def test_failed_legacy_repair_keeps_original_active_record(self):
        with self.database() as db:
            db.add(DocumentRegistry(
                topic_key='admin/1770000000 guide', source=str(self.root / '1770000000_guide_v1.docx'),
                version_label='v1', version_rank=1, sha256='a' * 64,
                generation_id='legacy', status='active',
            ))
        result = ingestion.repair_legacy_documents()
        self.assertEqual(result['documents'][0]['status'], 'failed')
        self.assertEqual(self.active()[0].generation_id, 'legacy')


class GroundingTests(unittest.TestCase):
    def service(self, documents=None, answer='Use VLAN 120. [S1]'):
        service = rag.RagService.__new__(rag.RagService)
        service.s = Settings(_env_file=None, database_url='sqlite://')
        service.prompt_budget = 1900
        service._active_generations = MagicMock(return_value={'current'})
        service.store = MagicMock()
        service.store.similarity_search_with_relevance_scores.return_value = documents or []
        service.llm = MagicMock()
        service.llm.invoke.return_value = SimpleNamespace(content=answer)
        return service

    def evidence(self, text='The approved VLAN is 120.', generation='current', score=.9):
        return (Document(page_content=text, metadata={
            'generation_id': generation, 'source': 'guide.pdf', 'version': 'v2', 'page': 4,
        }), score)

    def test_no_evidence_never_calls_model(self):
        service = self.service()
        result = service.answer('How do I configure it?', [])
        self.assertEqual(result['grounding_status'], 'insufficient_evidence')
        self.assertEqual(result['sources'], [])
        service.llm.invoke.assert_not_called()

    def test_empty_registry_does_not_query_vector_store(self):
        service = self.service()
        service._active_generations.return_value = set()
        service.answer('What VLAN?', [])
        service.store.similarity_search_with_relevance_scores.assert_not_called()
        service.llm.invoke.assert_not_called()

    def test_stale_or_low_relevance_evidence_is_excluded(self):
        service = self.service([self.evidence(generation='retired'), self.evidence(score=.1)])
        service.answer('What VLAN?', [])
        service.llm.invoke.assert_not_called()

    def test_only_registered_generation_is_queried_and_sources_match_citations(self):
        service = self.service([self.evidence()])
        result = service.answer('What VLAN?', [])
        self.assertEqual(result['sources'][0]['id'], 'S1')
        self.assertEqual(result['sources'][0]['page'], 4)
        query = service.store.similarity_search_with_relevance_scores.call_args
        self.assertEqual(query.kwargs['filter'], {'generation_id': {'$in': ['current']}})

    def test_unknown_or_missing_citations_return_deterministic_response(self):
        for answer in ('Use VLAN 999. [S99]', 'Use VLAN 999 without citations.'):
            service = self.service([self.evidence()], answer)
            result = service.answer('What VLAN?', [])
            self.assertEqual(result['grounding_status'], 'invalid_citations')
            self.assertNotIn('999', result['answer'])
            self.assertEqual(result['sources'], [])

    def test_followup_uses_previous_user_question_for_retrieval(self):
        service = self.service([self.evidence()])
        service.answer('Which port should it use?', [
            {'role': 'user', 'content': 'Configure switch model EX-123 VLAN 120'},
            {'role': 'assistant', 'content': 'An earlier model answer is not evidence.'},
        ])
        query = service.store.similarity_search_with_relevance_scores.call_args.args[0]
        self.assertIn('EX-123', query)
        self.assertIn('Which port', query)
        self.assertNotIn('earlier model answer', query)

    def test_combined_unicode_prompt_is_bounded(self):
        service = self.service([self.evidence('תיעוד מקומי ' * 10000)])
        service.answer('מה ההגדרה?', [
            {'role': 'user', 'content': 'שאלה קודמת ' * 10000},
            {'role': 'assistant', 'content': 'תשובה קודמת ' * 10000},
        ])
        prompt = service.llm.invoke.call_args.args[0]
        byte_count = sum(len(message.content.encode()) for message in prompt.messages)
        self.assertLessEqual(byte_count, service.prompt_budget)
        self.assertIn('[S1]', prompt.messages[1].content)

    def test_hebrew_no_evidence_response_is_localized(self):
        result = self.service().answer('איזה מתג מתאים?', [])
        self.assertIn('מסמכים המקומיים', result['answer'])


class SyncAndCacheTests(TemporaryStore):
    def test_failed_sync_job_retries_digest_and_preserves_remote_version(self):
        settings = self.settings
        settings.literature_site_url = 'http://intranet/manuals/'
        queue = MagicMock()
        queue.enqueue.side_effect = ['job-one', 'job-two']
        queue.status.return_value = None
        real_client = httpx.Client
        transport = httpx.MockTransport(lambda request: httpx.Response(200, content=b'Word bytes', request=request))
        with patch.object(literature, 'get_settings', return_value=settings):
            with patch.object(literature, 'JobQueue', return_value=queue):
                sync = literature.LiteratureSync()
        sync.discover = MagicMock(return_value=[
            literature.Remote('http://intranet/manuals/v7/manual.docx', 'manual.docx', 'remote/manual', 7, 'v7'),
        ])
        sync._registry_indexed = MagicMock(return_value=False)
        with patch.object(literature.httpx, 'Client', side_effect=lambda **kwargs: real_client(transport=transport, **kwargs)):
            first = sync.run_once()
            state = json.loads(sync.state.read_text())['files']['remote/manual']
            self.assertEqual(first['queued'], 1)
            self.assertEqual(state['status'], 'queued')
            queue.status.return_value = {'status': 'failed', 'error': 'Embedding unavailable'}
            second = sync.run_once()
            self.assertEqual(second['queued'], 1)
            calls = queue.enqueue.call_args_list
            self.assertNotEqual(calls[0].kwargs['idempotency_key'], calls[1].kwargs['idempotency_key'])
            payload = calls[1].args[2]
            self.assertEqual(payload['source_name'], 'manual.docx')
            self.assertEqual(payload['topic_key'], 'remote/manual')
            self.assertEqual(payload['version_rank'], 7)
            queue.status.return_value = {'status': 'succeeded', 'result': {'status': 'indexed'}}
            sync._registry_indexed.return_value = True
            third = sync.run_once()
            self.assertEqual(third['queued'], 0)
            self.assertEqual(third['indexed'], 1)

    def test_discovery_preserves_remote_directory_identity_and_query(self):
        self.settings.literature_site_url = 'http://intranet/manuals/'
        content = '<a href="/manuals/routers/v4/manual.pdf?download=1">manual.pdf</a>'
        real_client = httpx.Client
        transport = httpx.MockTransport(lambda request: httpx.Response(
            200, text=content, headers={'content-type': 'text/html'}, request=request,
        ))
        with patch.object(literature, 'get_settings', return_value=self.settings):
            with patch.object(literature, 'JobQueue', return_value=MagicMock()):
                sync = literature.LiteratureSync()
        with patch.object(literature.httpx, 'Client', side_effect=lambda **kwargs: real_client(transport=transport, **kwargs)):
            found = sync.discover()
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0].topic, 'internal_site/intranet/manuals/routers/manual')
        self.assertEqual(found[0].rank, 4)
        self.assertIn('?download=1', found[0].url)

    def test_internal_authentication_is_not_sent_to_external_redirect(self):
        self.settings.literature_site_url = 'http://intranet/manuals/'
        real_client = httpx.Client
        requests = []

        def redirect(request):
            requests.append(str(request.url))
            return httpx.Response(302, headers={'location': 'http://external.invalid/private'}, request=request)

        with patch.object(literature, 'get_settings', return_value=self.settings):
            with patch.object(literature, 'JobQueue', return_value=MagicMock()):
                sync = literature.LiteratureSync()
        with real_client(transport=httpx.MockTransport(redirect)) as client:
            with self.assertRaisesRegex(ValueError, 'outside'):
                sync._get(client, self.settings.literature_site_url)
        self.assertEqual(requests, ['http://intranet/manuals/'])

    def test_vision_cache_never_uses_default_embeddings(self):
        client = MagicMock()
        client.get_or_create_collection.return_value.get.return_value = {'documents': []}
        self.settings.vision_min_width = self.settings.vision_min_height = 1
        image = io.BytesIO()
        Image.new('RGB', (10, 10), 'white').save(image, 'PNG')
        response = MagicMock()
        response.json.return_value = {'response': 'Port 7'}
        with patch.object(vision, 'get_settings', return_value=self.settings):
            analyzer = vision.VisionAnalyzer(client)
        with patch.object(vision.httpx, 'Client') as http:
            http.return_value.__enter__.return_value.post.return_value = response
            self.assertEqual(analyzer.analyze(image.getvalue()), 'Port 7')
        self.assertIsNone(client.get_or_create_collection.call_args.kwargs['embedding_function'])
        self.assertEqual(analyzer.cache.upsert.call_args.kwargs['embeddings'], [[0.0]])

    def test_version_path_root_and_windows_separators_are_stable(self):
        identity = identify_version('manual_v2.pdf', r'\manuals\routers\version-4')
        self.assertEqual(identity.topic_key, 'manuals/routers/manual')
        self.assertEqual(identity.version_rank, 4)


if __name__ == '__main__':
    unittest.main()
