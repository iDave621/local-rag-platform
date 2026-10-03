"""Run against the live isolated stack; creates and removes its own test account/documents.

docker compose run --rm --no-deps -v <project>:/release:ro api python /release/tests/offline_acceptance.py
"""
import io
import json
import re
import secrets
import shutil
import sys
import tempfile
import time
import uuid
from pathlib import Path

sys.path.insert(0, '/app')

import fitz
import httpx
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from sqlalchemy import delete, select

from app.config import get_settings
from app.db import (AuditEvent, ChatSession, DocumentRegistry, JobRecord, User,
                    db_session)
from app.security import hash_password


def main():
    marker = uuid.uuid4().hex[:10]
    username = 'acceptance-' + marker
    password = secrets.token_urlsafe(24)
    checks = {}
    files_to_remove = []
    document_ids = []
    root = get_settings().data_path.resolve()
    with db_session() as db:
        user = User(username=username, password_hash=hash_password(password), role='admin')
        db.add(user); db.flush(); user_id = user.id
    # Test-only certificate bypass: the service DNS name differs from localhost SAN.
    # Application TLS and cookie settings remain enabled.
    client = httpx.Client(base_url='https://frontend', verify=False, timeout=30, trust_env=False)

    def record(name, condition):
        if not condition:
            raise AssertionError(name)
        checks[name] = True
        print('PASS', name, flush=True)

    def wait(jid, timeout=240):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            response = client.get('/api/jobs/' + jid)
            response.raise_for_status()
            row = response.json()
            if row['status'] in {'succeeded', 'failed'}:
                return row
            time.sleep(.5)
        raise AssertionError('Job did not finish within acceptance timeout')

    def upload(name, content):
        response = client.post('/api/admin/documents', files={'file': (name, content)})
        response.raise_for_status()
        jid = response.json()['job_id']
        with db_session() as db:
            files_to_remove.append(Path(db.get(JobRecord, jid).payload['path']))
        return wait(jid)

    def ask(message, sid=None):
        if sid is None:
            response = client.post('/api/sessions'); response.raise_for_status()
            sid = response.json()['session_id']
        response = client.post('/api/chat', json={'session_id': sid, 'message': message,
                               'request_id': str(uuid.uuid4())})
        record('async_chat_response', response.status_code == 202)
        row = wait(response.json()['job_id'])
        if row['status'] != 'succeeded':
            raise AssertionError(row['error'])
        return row['result'], sid

    try:
        response = client.post('/api/auth/login', json={'username': username, 'password': password})
        response.raise_for_status()
        record('secure_login_cookie', 'Secure' in response.headers['set-cookie'] and 'HttpOnly' in response.headers['set-cookie'])
        response = client.get('/api/health')
        record('all_dependencies_ready', response.status_code == 200)
        with httpx.Client(timeout=3, trust_env=False) as blocked:
            try:
                blocked.get('https://example.com')
            except httpx.HTTPError:
                record('backend_external_network_blocked', True)
            else:
                raise AssertionError('Run acceptance with the offline-validation Compose overlay')

        for version, vlan in [(1, 321), (2, 654)]:
            doc = Document()
            doc.add_paragraph(f'The Calibration Gateway {marker} approved management configuration follows.')
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = f'Approved management VLAN for Calibration Gateway {marker}'
            table.cell(0, 1).text = str(vlan)
            data = io.BytesIO(); doc.save(data)
            row = upload(f'Acceptance_{marker}_v{version}.docx', data.getvalue())
            record(f'word_v{version}_table_indexed', row['status'] == 'succeeded' and row['result']['status'] == 'indexed')
        with db_session() as db:
            docs = db.scalars(select(DocumentRegistry).where(DocumentRegistry.source.like(f'Acceptance_{marker}%'))).all()
            active = [row for row in docs if row.status == 'active']
            record('one_current_document_version', len(active) == 1 and active[0].version_rank == 2)
            document_ids.extend(row.id for row in active)
        answer, sid = ask(f'What is the approved management VLAN for Calibration Gateway {marker}?')
        record('cited_word_table_answer', '654' in answer['answer'] and any(
            source['source'] == f'Acceptance_{marker}_v2.docx' for source in answer['sources']))
        history = client.get(f'/api/sessions/{sid}/history').json()
        record('history_retains_citations', len(history) == 2 and bool(history[-1]['sources']))
        record('history_session_restored', any(row['session_id'] == sid for row in client.get('/api/sessions').json()))

        pdf = fitz.open()
        for _ in range(2):
            page = pdf.new_page()
            page.insert_text((72, 72), f'Acceptance pump {marker} emergency reset procedure: press RESET for 9 seconds.')
        row = upload(f'AcceptancePDF_{marker}_v1.pdf', pdf.tobytes())
        pdf.close()
        record('repeated_pdf_pages_indexed', row['status'] == 'succeeded' and row['result']['chunks'] == 2)

        image = Image.new('RGB', (1500, 300), 'white')
        font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 40)
        ImageDraw.Draw(image).text((30, 60), 'ScanGate approved management VLAN is 777.', fill='black', font=font)
        png = io.BytesIO(); image.save(png, format='PNG')
        pdf = fitz.open(); page = pdf.new_page(width=750, height=150)
        page.insert_image(page.rect, stream=png.getvalue())
        row = upload(f'AcceptanceScan_{marker}_v1.pdf', pdf.tobytes()); pdf.close()
        record('scanned_pdf_ocr_indexed', row['status'] == 'succeeded' and row['result']['chunks'] > 0)
        answer, _ = ask('What is the approved management VLAN for ScanGate?')
        record('cited_scanned_pdf_answer', '777' in answer['answer'] and any(
            source['source'] == f'AcceptanceScan_{marker}_v1.pdf' and source['page'] == 1
            for source in answer['sources']))

        answer, hebrew_sid = ask('מהו VLAN הניהול המאושר עבור ScanGate?')
        record('cited_hebrew_answer', '777' in answer['answer'] and bool(
            re.search(r'[\u0590-\u05ff]', answer['answer'])) and any(
            source['source'] == f'AcceptanceScan_{marker}_v1.pdf' and source['page'] == 1
            for source in answer['sources']))
        answer, _ = ask('ומה לגבי אותו VLAN?', hebrew_sid)
        record('followup_retains_document_context', '777' in answer['answer'] and any(
            source['source'] == f'AcceptanceScan_{marker}_v1.pdf'
            for source in answer['sources']))

        doc = Document(); data = io.BytesIO(); doc.save(data)
        row = upload(f'AcceptanceEmpty_{marker}.docx', data.getvalue())
        record('empty_document_failure_visible', row['status'] == 'failed' and 'No readable text' in row['error'])
        answer, _ = ask('What is the fusion temperature of fictional reactor ZXQ991 on planet Nonsense?')
        record('unknown_question_refused', not answer['sources'])

        listing = client.get('/api/admin/documents').json()
        document_ids = [row['id'] for row in listing if marker in row['source_name']]
        for document_id in document_ids:
            record('document_removal', client.delete(f'/api/admin/documents/{document_id}').status_code == 204)
        with db_session() as db:
            record('removed_documents_unsearchable', not db.scalar(select(DocumentRegistry.id).where(
                DocumentRegistry.source.like(f'%{marker}%'), DocumentRegistry.status == 'active')))
        print(json.dumps({'ok': True, 'checks': checks}, indent=2), flush=True)
    finally:
        # Retire only records whose original source contains this random test marker.
        from app.rag import RagService
        with db_session() as db:
            rows = db.scalars(select(DocumentRegistry).where(DocumentRegistry.source.like(f'%{marker}%'))).all()
            generations = [row.generation_id for row in rows]
            for row in rows:
                row.status = 'deleted'
        if generations:
            rag = RagService()
            for generation in generations:
                rag.delete_generation(generation)
        with db_session() as db:
            db.execute(delete(DocumentRegistry).where(DocumentRegistry.source.like(f'%{marker}%')))
            db.execute(delete(JobRecord).where(JobRecord.user_id == user_id))
            db.execute(delete(AuditEvent).where(AuditEvent.user_id == user_id))
            db.execute(delete(User).where(User.id == user_id))
        for path in files_to_remove:
            resolved = path.resolve()
            if resolved.is_relative_to(root / 'staging' / 'admin') and resolved.parent != root / 'staging' / 'admin':
                resolved.unlink(missing_ok=True)
                if resolved.parent.is_dir() and not any(resolved.parent.iterdir()):
                    resolved.parent.rmdir()
        client.close()


if __name__ == '__main__':
    main()
